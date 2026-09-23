# -*- coding: utf-8 -*-
"""Thử XUẤT BÁO CÁO trên Chromium thật (Playwright): mọi màn · nút Excel và nút PDF trên thanh đầu trang.

    <python có playwright + openpyxl> kiem/thu_xuat_bao_cao.py [http://127.0.0.1:8011]

Mỗi màn của Sếp (xem hết) và của Bãi (không thấy tiền):
  · bấm Excel → bắt tệp tải về → MỞ LẠI bằng openpyxl: đúng .xlsx, có đầu trang, số dòng dữ liệu = số dòng bảng đang
    hiện, ô tiền là SỐ (không phải chữ) và mang tiền tệ trong định dạng ô;
  · Bãi: tệp không được chứa cột tiền mà màn ẩn với Bãi (chữ "Thành tiền", "Doanh thu", "Đơn giá"…);
  · bấm PDF (giữ bản in) → in ra tệp PDF thật, có đầu trang báo cáo; tờ phiếu thì giữ khổ in phiếu.
Không ghi gì vào máy chủ.
"""
import io
import os
import sys
import tempfile

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

GOC = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8011"
RA = os.path.join(tempfile.gettempdir(), "epl_thu_xuat")
os.makedirs(RA, exist_ok=True)
CHU_TIEN_BAN = ("Doanh thu", "Cước", "Lãi", "Thành tiền", "Đơn giá", "Giá vốn")    # Bãi không được thấy
DEM_DONG = """() => { const r = document.querySelector('#noi-dung .mod-root'); const hien = e => e.getClientRects().length > 0;
  return [...r.querySelectorAll('table')].filter(t => hien(t) && !t.closest('dialog,.hop-thoai') && !t.parentElement.closest('table'))
    .map(t => [...t.tBodies, ...(t.tFoot ? [t.tFoot] : [])].flatMap(b => [...b.rows]).filter(tr => hien(tr) && !tr.querySelector('td.empty') && [...tr.cells].some(c => hien(c) && c.innerText.trim())).length); }"""


def dang_nhap(pg, u):
    pg.goto(GOC + "/", wait_until="networkidle")
    pg.fill("#lgU", u); pg.fill("#lgP", "1234"); pg.click("#lgBtn")
    pg.wait_for_selector("#app:not([hidden])", timeout=15000); pg.wait_for_timeout(600)
    if pg.eval_on_selector("#langApp .ln-nut", "b => b.title") != "Tiếng Việt":
        pg.click("#langApp .ln-nut"); pg.wait_for_timeout(250); pg.click('#langApp .ln-muc[data-lang="vi"]'); pg.wait_for_timeout(400)


def kiem_excel(pg, u, m, loi):
    dong_man = pg.evaluate(DEM_DONG)
    tu_dung = pg.evaluate("m => !!(EPL.modules[m] && EPL.modules[m].xuatExcel)", m)
    if not sum(dong_man) and not tu_dung:
        pg.click("#btnExcel"); pg.wait_for_timeout(300)
        t = pg.inner_text("body")
        return "không bảng → báo 'chưa có bảng'" if "chưa có bảng" in t else "không bảng (KHÔNG báo gì)"
    with pg.expect_download(timeout=15000) as d:
        pg.click("#btnExcel")
    tep = os.path.join(RA, "%s_%s.xlsx" % (u, m)); d.value.save_as(tep)
    wb = load_workbook(tep)                                   # mở được = đúng định dạng .xlsx
    so_dong, so_tien, chu_tien, tien_chu = [], 0, [], []
    for ws in wb.worksheets:
        dau = ws["A1"].value or ""
        if not str(dau).startswith("EPL"):
            loi.append("%s/%s: sheet %s thiếu đầu trang" % (u, m, ws.title))
        # dòng đầu cột = dòng đầu tiên có ô tô nền (style header) sau khối đầu trang
        hang = list(ws.iter_rows())
        i_dau = next((i for i, r in enumerate(hang) if any(c.fill and c.fill.fgColor and c.fill.fgColor.rgb == "FFE3EDE7" for c in r)), None)
        if i_dau is None:
            so_dong.append(0); continue
        j = i_dau
        while j + 1 < len(hang) and any(c.fill and c.fill.fgColor and c.fill.fgColor.rgb == "FFE3EDE7" for c in hang[j + 1]):
            j += 1
        than = [r for r in hang[j + 1:] if any(c.value not in (None, "") for c in r)]
        so_dong.append(len(than))
        for r in hang[i_dau:j + 1]:
            for c in r:
                if isinstance(c.value, str) and u == "thabok" and any(c.value.startswith(k) for k in CHU_TIEN_BAN):
                    chu_tien.append(c.value)
        for r in than:
            for c in r:
                if isinstance(c.value, (int, float)):
                    if any(k in (c.number_format or "") for k in ('"LAK"', '"USD"', '"THB"', '"VND"', '"CNY"')):
                        so_tien += 1
                elif isinstance(c.value, str) and __import__("re").match(r"^-?[\d,]+(\.\d+)? (LAK|USD|THB|VND|CNY)$", c.value):
                    tien_chu.append(c.value)
    if chu_tien:
        loi.append("%s/%s: Bãi xuất ra cột tiền %s" % (u, m, chu_tien[:3]))
    if tien_chu:
        loi.append("%s/%s: ô tiền còn là CHỮ, Excel không cộng được: %s" % (u, m, tien_chu[:3]))
    man = [n for n in dong_man if n]
    if not tu_dung and sorted(n for n in so_dong if n) != sorted(man):
        loi.append("%s/%s: số dòng Excel %s ≠ màn hình %s" % (u, m, so_dong, dong_man))
    return "%d sheet · %s dòng · %d ô tiền là số" % (len(wb.worksheets), "+".join(map(str, so_dong)), so_tien)


def kiem_pdf(pg, u, m, loi):
    giay = pg.evaluate("() => { window.__don = EPL.xuatPDF(true); return !!document.querySelector('.in-bao-cao') ; }")
    co_dau = pg.evaluate("() => !!document.querySelector('.in-dau')")
    tep = os.path.join(RA, "%s_%s.pdf" % (u, m))
    pg.emulate_media(media="print")
    pg.pdf(path=tep, prefer_css_page_size=True, print_background=True)
    pg.emulate_media(media="screen")
    pg.evaluate("() => window.__don && window.__don()")
    if not os.path.getsize(tep) > 2000:
        loi.append("%s/%s: PDF rỗng" % (u, m))
    if giay and not co_dau:
        loi.append("%s/%s: PDF báo cáo thiếu đầu trang" % (u, m))
    if pg.evaluate("() => !!document.querySelector('.in-dau') || document.body.classList.contains('in-bao-cao')"):
        loi.append("%s/%s: in xong không dọn đầu trang" % (u, m))
    return ("báo cáo ngang" if giay else "tờ phiếu") + " · %d KB" % (os.path.getsize(tep) // 1024)


def main():
    loi, n = [], 0
    with sync_playwright() as p:
        b = p.chromium.launch()
        for u in ("admin", "thabok"):
            ctx = b.new_context(viewport={"width": 1440, "height": 900}, locale="vi-VN", accept_downloads=True)
            pg = ctx.new_page()
            pg.on("pageerror", lambda e, u=u: loi.append("%s JS: %s" % (u, e)))
            dang_nhap(pg, u)
            mods = pg.eval_on_selector_all("#nav [data-mod]", "bs => [...new Set(bs.map(b => b.dataset.mod))]")
            for m in mods:
                pg.evaluate("h => { location.hash = h; }", "#/" + m); pg.wait_for_timeout(1700)
                if pg.is_hidden("#xuatNut"):
                    print("%-8s %-16s (không có nút xuất)" % (u, m)); continue
                n += 1
                try:
                    e = kiem_excel(pg, u, m, loi)
                    f = kiem_pdf(pg, u, m, loi)
                except Exception as ex:  # noqa: BLE001
                    loi.append("%s/%s: %s" % (u, m, str(ex)[:200])); e = f = "HỎNG"
                print("%-8s %-16s Excel: %-40s PDF: %s" % (u, m, e, f), flush=True)
            ctx.close()
        b.close()
    print("\nTệp thử để mở tay: %s" % RA)
    if loi:
        print("\nHỎNG — %d chỗ:" % len(loi))
        for x in loi:
            print("  ·", x)
        sys.exit(1)
    print("✅ XUẤT BÁO CÁO: %d màn · Excel mở lại được, số dòng khớp màn hình, tiền là số · PDF có đầu trang · Bãi không lộ tiền" % n)


if __name__ == "__main__":
    main()
